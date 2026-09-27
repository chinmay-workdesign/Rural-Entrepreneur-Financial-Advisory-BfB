"""
Admin analytics page.

1. Counts come from the database: contacts per channel, the funnel, DPRs and districts
2. Filtering to one district shows only that district's people and applications
3. Only administrators can open the page or the API; officers are sent back, anonymous users to /login
"""
import uuid

from fastapi.testclient import TestClient

from app import analytics
from app.db import crud
from app.db.models import Beneficiary, EnterpriseProposal
from app.db.session import SessionLocal
from app.main import app
from tests.intake_helpers import login_officer


def _district() -> str:
    return "Testpur" + uuid.uuid4().hex[:6]


def _seed(district: str):
    """Three people in `district`: one only picked a language, one is answering questions, one has a DPR."""
    db = SessionLocal()
    try:
        tag = uuid.uuid4().hex[:8]
        people = [
            Beneficiary(telegram_chat_id=f"a{tag}", primary_channel="telegram", conversation_state="LANGUAGE_SELECTION",
                        district=district, preferred_language="kannada"),
            Beneficiary(whatsapp_number=f"91{tag}1", primary_channel="whatsapp", conversation_state="COLLECTING",
                        district=district.lower(), preferred_language="hindi"),
            Beneficiary(whatsapp_number=f"91{tag}2", primary_channel="whatsapp", conversation_state="SUBMITTED",
                        district=district, preferred_language="kannada"),
        ]
        db.add_all(people)
        db.commit()
        p = EnterpriseProposal(beneficiary_id=people[2].id, business_trade="Dairy unit", scheme_tier="TERM_LOAN",
                               project_cost=200000, sanctioned_loan=180000, beneficiary_margin=20000, monthly_emi=0,
                               status="SANCTIONED", dpr_pdf_url="/static/dprs/x.pdf")
        db.add(p)
        db.commit()
        crud.save_proposal_snapshot(db, p.id, {"district": district, "profile": {
            "gender": "female", "social_category": "sc", "age": 34, "area_type": "rural"},
            "financial_structure": {"agency": "NSFDC"}})
    finally:
        db.close()


def _admin() -> TestClient:
    client = TestClient(app)
    db = SessionLocal()
    try:
        crud.seed_default_users(db)
    finally:
        db.close()
    assert client.post("/auth/login", json={"email": "admin@sca.gov.in", "password": "Admin@123"}).status_code == 200
    return client


# 1 + 2. Counts and the district filter
def test_district_statistics_are_counted_from_the_database():
    district = _district()
    _seed(district)
    _seed(_district())  # another district that must not leak into the filtered view

    db = SessionLocal()
    try:
        everything = analytics.compute(db)
        one = analytics.compute(db, district=district.upper())
    finally:
        db.close()

    row = next(r for r in everything["districts"] if r["district"] == district.title())
    assert (row["contacts"], row["people_with_dpr"], row["sanctioned"], row["project_cost"]) == (3, 1, 1, 200000.0)

    assert one["district"] == district.title()
    s = one["summary"]
    assert (s["contacts"], s["contacts_telegram"], s["contacts_whatsapp"]) == (3, 1, 2)
    assert (s["dprs_generated"], s["sanctioned"], s["total_loan_sanctioned"], s["conversion_pct"]) == (1, 1, 180000.0, 33.3)
    assert [f["count"] for f in one["funnel"]] == [3, 2, 1, 1, 1, 1, 1]
    assert {i["label"]: i["count"] for i in one["languages"]} == {"Kannada": 2, "Hindi": 1}
    assert one["gender"] == [{"label": "Women", "count": 1}]
    assert one["schemes"] == [{"label": "NSFDC", "count": 1}]
    assert one["age"] == [{"label": "26–35", "count": 1}]
    assert len(one["trend"]) == analytics.TREND_DAYS and sum(t["contacts"] for t in one["trend"]) == 3


# 3. Access
def test_analytics_is_admin_only():
    anon = TestClient(app)
    assert anon.get("/internal/analytics").status_code == 401
    page = anon.get("/admin/analytics", follow_redirects=False)
    assert page.status_code == 302 and page.headers["location"] == "/login"

    officer = TestClient(app)
    login_officer(officer)
    assert officer.get("/internal/analytics").status_code == 403
    page = officer.get("/admin/analytics", follow_redirects=False)
    assert page.status_code == 302 and page.headers["location"] == "/admin"

    admin = _admin()
    assert admin.get("/admin/analytics").status_code == 200
    district = _district()
    _seed(district)
    body = admin.get("/internal/analytics", params={"district": district}).json()
    assert body["summary"]["contacts"] == 3 and body["district"] == district.title()


# 4. Excel export
def test_excel_export_has_every_sheet_and_respects_the_district():
    import io
    from openpyxl import load_workbook

    district = _district()
    _seed(district)
    db = SessionLocal()
    try:  # applicant text that looks like a formula must stay plain text
        b = db.query(Beneficiary).filter(Beneficiary.district == district.lower()).first()
        b.full_name = '=HYPERLINK("http://x","click")'
        db.commit()
    finally:
        db.close()

    officer = TestClient(app)
    login_officer(officer)
    assert officer.get("/internal/analytics/export").status_code == 403
    assert TestClient(app).get("/internal/analytics/export").status_code == 401

    res = _admin().get("/internal/analytics/export", params={"district": district})
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert f"advisor-analytics-{district.lower()}-" in res.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Summary", "Funnel", "Districts", "Breakdowns", "Last 30 days", "Applications", "People contacted"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=7, values_only=True)}
    assert wb["Summary"]["B2"].value == district.title()
    assert (summary["People who contacted the bot"], summary["DPRs generated"], summary["Total loans sanctioned"]) == (3, 1, 180000)

    apps = list(wb["Applications"].iter_rows(values_only=True))
    assert len(apps) == 2 and apps[1][apps[0].index("Status")] == "Sanctioned"
    assert apps[1][apps[0].index("Corporation")] == "NSFDC" and apps[1][apps[0].index("Project cost (Rs.)")] == 200000

    people = list(wb["People contacted"].iter_rows(values_only=True))
    assert len(people) == 4 and {r[people[0].index("District")] for r in people[1:]} == {district.title()}
    name_cell = next(c for c in wb["People contacted"]["E"] if c.value and "HYPERLINK" in str(c.value))
    assert name_cell.data_type == "s"

    # Every district is listed on the Districts sheet, whatever the scope
    assert district.title() in [r[0] for r in wb["Districts"].iter_rows(min_row=2, values_only=True)]
