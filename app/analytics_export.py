"""
Excel export of the admin analytics page: the same statistics (overall or one district) plus the full list of
applications and of people who contacted the bot.
"""
import io
from datetime import datetime
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import analytics

_HEAD_FILL = PatternFill("solid", fgColor="1E5B42")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14, color="10251D")
_SECTION_FONT = Font(bold=True, size=11, color="1E5B42")
_MONEY = '"Rs. "#,##0'
_DATE = "dd-mmm-yyyy"
_DATETIME = "dd-mmm-yyyy hh:mm"
_PCT = '0.0"%"'


def _put(ws, row: int, col: int, value: Any, fmt: Optional[str] = None, font: Optional[Font] = None):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # applicant text is never read as a formula, even if it starts with "="
    if fmt:
        cell.number_format = fmt
    if font:
        cell.font = font
    return cell


def _header(ws, row: int, labels: List[str]):
    for i, label in enumerate(labels, 1):
        c = _put(ws, row, i, label, font=_HEAD_FONT)
        c.fill = _HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _fmt_for(key: str, value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return _DATETIME
    if "(Rs.)" in key or key in ("Project cost requested",):
        return _MONEY
    return None


def _table(ws, rows: List[Dict[str, Any]], empty: str):
    if not rows:
        _put(ws, 1, 1, empty)
        return
    keys = list(rows[0].keys())
    _header(ws, 1, keys)
    for r, row in enumerate(rows, 2):
        for c, k in enumerate(keys, 1):
            _put(ws, r, c, row[k], _fmt_for(k, row[k]))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{len(rows) + 1}"


def _autosize(ws, max_width: int = 45):
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        if any(isinstance(c.value, datetime) for c in col):
            width = max(width, 18)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), max_width)


def build_workbook(db: Session, district: Optional[str] = None) -> bytes:
    data = analytics.compute(db, district=district)
    scope = data["district"] or "All districts"
    s = data["summary"]
    wb = Workbook()

    # 1. Summary
    ws = wb.active
    ws.title = "Summary"
    _put(ws, 1, 1, "Rural Enterprise Advisor: analytics", font=_TITLE_FONT)
    _put(ws, 2, 1, "Scope")
    _put(ws, 2, 2, scope, font=Font(bold=True))
    _put(ws, 3, 1, "Generated (IST)")
    _put(ws, 3, 2, datetime.fromisoformat(data["generated_at"]).replace(tzinfo=None), _DATETIME)
    _put(ws, 4, 1, "Every figure is counted from the application database; values applicants did not give are "
                   "shown as 'Not stated'.")
    metrics = [
        ("People who contacted the bot", s["contacts"], None),
        ("  on Telegram", s["contacts_telegram"], None),
        ("  on WhatsApp", s["contacts_whatsapp"], None),
        ("DPRs generated", s["dprs_generated"], None),
        ("People who generated a DPR", s["people_with_dpr"], None),
        ("Contact to DPR conversion", s["conversion_pct"], _PCT),
        ("Applications", s["applications"], None),
        ("Awaiting review", s["awaiting"], None),
        ("Sent back", s["sent_back"], None),
        ("Sanctioned", s["sanctioned"], None),
        ("Rejected", s["rejected"], None),
        ("Field verified", s["field_verified"], None),
        ("Total project cost requested", s["total_project_cost"], _MONEY),
        ("Average project cost per application", s["avg_project_cost"], _MONEY),
        ("Total loans requested", s["total_loan_requested"], _MONEY),
        ("Applicants' own contribution", s["total_own_contribution"], _MONEY),
        ("Total loans sanctioned", s["total_loan_sanctioned"], _MONEY),
    ]
    _header(ws, 6, ["Measure", "Value"])
    for i, (label, value, fmt) in enumerate(metrics, 7):
        _put(ws, i, 1, label)
        _put(ws, i, 2, value, fmt)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 22

    # 2. Funnel
    ws = wb.create_sheet("Funnel")
    _header(ws, 1, ["Step", "People", "% of people contacted"])
    top = data["funnel"][0]["count"]
    for i, f in enumerate(data["funnel"], 2):
        _put(ws, i, 1, f"{i - 1}. {f['label']}")
        _put(ws, i, 2, f["count"])
        _put(ws, i, 3, round(100 * f["count"] / top, 1) if top else None, _PCT)
    _put(ws, len(data["funnel"]) + 3, 1, "A person counts for a step if their current conversation reached it or "
                                         "they have an application on file (/start clears earlier answers).")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 22

    # 3. Districts (always every district)
    ws = wb.create_sheet("Districts")
    _table(ws, [{"District": r["district"], "People contacted": r["contacts"],
                 "People who generated a DPR": r["people_with_dpr"], "Contact to DPR %": r["conversion_pct"],
                 "Applications": r["applications"], "DPRs generated": r["dprs"], "Awaiting review": r["awaiting"],
                 "Sanctioned": r["sanctioned"], "Rejected": r["rejected"], "Project cost requested": r["project_cost"]}
                for r in data["districts"]], "No contacts yet.")
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = _PCT
    _autosize(ws)

    # 4. Breakdowns
    ws = wb.create_sheet("Breakdowns")
    r = 1
    for key, title in (("current_stage", "Where people are now (people contacted)"),
                       ("channels", "Channel (people contacted)"), ("languages", "Language chosen (people contacted)"),
                       ("status", "Application status (applications)"), ("schemes", "Scheme / corporation (applications)"),
                       ("businesses", "Businesses, as named by applicants (applications)"),
                       ("gender", "Gender (applicants)"), ("social_category", "Social category (applicants)"),
                       ("age", "Age (applicants)"), ("area", "Area (applicants)")):
        _put(ws, r, 1, title, font=_SECTION_FONT)
        _header(ws, r + 1, ["", "Count", "%"])
        items = data[key]
        total = sum(i["count"] for i in items)
        for j, item in enumerate(items, r + 2):
            _put(ws, j, 1, item["label"])
            _put(ws, j, 2, item["count"])
            _put(ws, j, 3, round(100 * item["count"] / total, 1) if total else None, _PCT)
        if not items:
            _put(ws, r + 2, 1, "No data yet.")
        r += max(len(items), 1) + 3
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10

    # 5. Daily trend
    ws = wb.create_sheet("Last 30 days")
    _header(ws, 1, ["Date (IST)", "New contacts", "DPRs generated"])
    for i, t in enumerate(data["trend"], 2):
        _put(ws, i, 1, datetime.fromisoformat(t["day"]), _DATE)
        _put(ws, i, 2, t["contacts"])
        _put(ws, i, 3, t["dprs"])
    for col, w in (("A", 16), ("B", 14), ("C", 16)):
        ws.column_dimensions[col].width = w

    # 6 + 7. Full lists
    ws = wb.create_sheet("Applications")
    _table(ws, analytics.application_rows(db, district), "No applications yet.")
    _autosize(ws)
    ws = wb.create_sheet("People contacted")
    _table(ws, analytics.contact_rows(db, district), "Nobody has contacted the bot yet.")
    _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
